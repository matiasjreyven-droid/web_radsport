# -*- coding: utf-8 -*-
import openpyxl, json, re, sys, os, unicodedata, shutil
sys.stdout.reconfigure(encoding='utf-8')

PROD_DIR = r"C:\Users\matia\OneDrive\Escritorio\Radsport\06_Pagina Web\marcas\leovince\productos"
SRC = r"C:\Users\matia\OneDrive\Escritorio\Radsport\01_Importacion\01_Marcas\07_Leo Vince\Lista de Precios LEO VINCE - Abr 26.xlsx"

def clean(v): return str(v).strip() if v is not None else ""
def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii','ignore').decode().lower()
    s = re.sub(r'[^\w\s-]','',s); s = re.sub(r'[\s_]+','-',s).strip('-'); return s
def load_xlsx(src, local):
    try: return openpyxl.load_workbook(src, data_only=True)
    except PermissionError:
        try: shutil.copy(src, local)
        except PermissionError: pass
        return openpyxl.load_workbook(local, data_only=True)

import hashlib
from PIL import Image
import numpy as np
files = set(os.listdir(PROD_DIR))
def _md5(f):
    return hashlib.md5(open(os.path.join(PROD_DIR, f), 'rb').read()).hexdigest()
_metcache = {}
def _metrics(f):
    if f not in _metcache:
        try:
            a = np.asarray(Image.open(os.path.join(PROD_DIR, f)).convert('RGB'))
            m = a.min(axis=2) < 235
            nw = float(m.mean())
            if m.any():
                ys, xs = np.where(m); asp = (xs.max()-xs.min()+1) / (ys.max()-ys.min()+1)
            else:
                asp = 0.0
            _metcache[f] = (nw, float(asp))
        except Exception:
            _metcache[f] = (1.0, 0.0)
    return _metcache[f]
def _score(f):
    nw, asp = _metrics(f)
    return asp * (1.0 - nw)      # el escape: ancho (aspect alto) y limpio (poco contenido)

# --- Overrides validados visualmente (contact sheets lv_sheet_*.png) ---
# codigos donde el "-studio" es en realidad una MOTO y el escape es una foto numerada
ESC_SUFFIX = {'14352eb': '1', '14414ebu': '1'}
# codigos SIN foto de escape propia (solo motos/diagrama) -> portada = logo LV
# ojo: 80014/80045/80046 son el codigo real (los archivos llevan un 0 extra: 800140-...)
NO_ESCAPE  = {'14363eu', '14419ebu', '14421eu', '80014', '80045', '80046'}
# codigos sin escape que pueden tomar prestado el de un hermano
BORROW     = {'14419ebu': '14419e'}

def _prod_files(pref):
    studio = sorted(f for f in files if f.startswith(pref + '-studio'))
    nums = sorted((f for f in files if re.match(re.escape(pref) + r'-\d+\.', f)),
                  key=lambda x: int(re.search(r'-(\d+)\.', x).group(1)))
    bikes = sorted((f for f in files if re.match(re.escape(pref) + r'-bike-\d+\.', f)),
                   key=lambda x: int(re.search(r'-(\d+)\.', x).group(1)))
    return studio, nums, bikes

def photos_for(code):
    """Devuelve (escapes, mounted): escapes = foto(s) del escape solo (portada + galeria);
    mounted = fotos del escape puesto en la moto (van debajo, 'como queda instalado')."""
    cl = code.lower()
    for pref in [cl, cl + '0']:
        studio, nums, bikes = _prod_files(pref)
        if not (studio or nums or bikes):
            continue
        # dedup por contenido (studio==-1, bike-1==bike-2, etc.)
        seen = set(); uniq = []
        for f in studio + nums + bikes:
            hh = _md5(f)
            if hh not in seen: seen.add(hh); uniq.append(f)

        escapes, mounted = [], []
        forced = ESC_SUFFIX.get(cl)                    # sufijo del escape forzado
        for f in uniq:
            suf = f[len(pref) + 1:].rsplit('.', 1)[0]  # studio / 1 / bike-1 ...
            is_escape = (suf == forced) if forced else (suf == 'studio')
            (escapes if is_escape else mounted).append(f)

        if cl in NO_ESCAPE:                            # el "-studio" era una moto -> a mounted
            mounted = escapes + mounted; escapes = []
            if cl in BORROW:                           # tomar el escape de un hermano
                bs, bn, _ = _prod_files(BORROW[cl])
                if bs: escapes = [bs[0]]
        # el escape va primero por score (limpio+ancho) por si hubiera varios
        escapes.sort(key=_score, reverse=True)
        return escapes, mounted[:6]
    return [], []

def disp(f):   # usar la version recortada (productos/c/) para mostrar; original para el orden
    return 'productos/c/'+f if os.path.exists(os.path.join(PROD_DIR, 'c', f)) else 'productos/'+f

def material_of(modelo):
    m = modelo.lower()
    if 'negro' in m or 'black' in m: return 'Negro'
    if 'carbon' in m: return 'Carbono'
    if 'inoxidable' in m or 'inox' in m: return 'Inoxidable'
    return 'Otro'

# Modelos abreviados con "/" -> motos individuales
EXPAND = {
 'Z400/500': ['Z 400', 'Z 500'], 'Ninja 400/500': ['Ninja 400', 'Ninja 500'],
 '690 Enduro/R': ['690 Enduro', '690 Enduro R'], '690 SMC/R': ['690 SMC', '690 SMC R'],
 '790/890 Adventure': ['790 Adventure', '890 Adventure'],
 '1290 Super Adv R/S/T': ['1290 Super Adventure R', '1290 Super Adventure S', '1290 Super Adventure T'],
 '1190 Adv/R': ['1190 Adventure', '1190 Adventure R'], '1090 Adv/R': ['1090 Adventure', '1090 Adventure R'],
}
def expand_models(mm):
    mm = mm.strip()
    if mm in EXPAND: return EXPAND[mm]
    if ' / ' in mm: return [x.strip() for x in mm.split(' / ') if x.strip()]
    return [mm]

wb = load_xlsx(SRC, 'lv_precios.xlsx'); ws = wb['Leo Vince']
products = {}; motos = []; order = 0
for r in ws.iter_rows(min_row=7, values_only=True):
    code = clean(r[2])
    if not code: continue
    modelo_full = clean(r[3]); parts = [p.strip() for p in modelo_full.split('|')]
    name = parts[0]; material_txt = parts[1] if len(parts) > 1 else ''
    anios = clean(r[4]).replace('*','').strip()
    marca_moto = clean(r[5]); modelo_moto = clean(r[6]); url = clean(r[8])
    escph, mntph = photos_for(code)
    if code not in products:
        order += 1
        products[code] = dict(code=code, name=name, material=material_txt, material_cat=material_of(modelo_full),
                              anios=anios, url=url,
                              photos=[disp(f) for f in escph],               # escape solo (portada + galeria), recortadas
                              mounted=[disp(f) for f in mntph],              # escape puesto en la moto ("como queda instalado")
                              main=disp(escph[0]) if escph else '../../img/logos/leovince.png',
                              has_photo=bool(escph), motos=[], order=order)
    for mm_raw in [x.strip() for x in modelo_moto.split('|') if x.strip()]:
        for mm in expand_models(mm_raw):
            products[code]['motos'].append({'brand': marca_moto, 'model': mm, 'anios': anios})
            motos.append({'brand': marca_moto, 'model': mm, 'anios': anios, 'code': code, 'id': slug(marca_moto+'-'+mm)})

json.dump(dict(products=products, motos=motos), open('lv_data.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
print(f"Escapes: {len(products)} | Motos: {len(motos)} | con foto propia: {sum(1 for p in products.values() if p['has_photo'])}")
