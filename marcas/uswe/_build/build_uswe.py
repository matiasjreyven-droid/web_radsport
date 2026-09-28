# -*- coding: utf-8 -*-
import openpyxl, json, re, sys, os, unicodedata, shutil
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

BASE = Path(r"C:\Users\matia\OneDrive\Escritorio\Radsport\06_Pagina Web\marcas\uswe")
PRODdir = BASE / 'productos'
SRC = r"C:\Users\matia\OneDrive\Escritorio\Radsport\01_Importacion\01_Marcas\06_USWE-Giant Loop\01_Listas de precios\Lista de Precios USWE & GIANT LOOP - RG.xlsx"

def clean(v): return str(v).strip() if v is not None else ""
def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii','ignore').decode().lower()
    s = re.sub(r'[^\w\s/-]','',s); s = re.sub(r'[\s/_]+','-',s).strip('-'); s = re.sub(r'-+','-',s); return s
def load_xlsx(src, local):
    try: return openpyxl.load_workbook(src, data_only=True)
    except PermissionError:
        try: shutil.copy(src, local)
        except PermissionError: pass
        return openpyxl.load_workbook(local, data_only=True)

files = os.listdir(PRODdir)

# --- 1) catalogo actual: id -> {category, segment, name, vol, color->photo, main} ---
cat = (BASE / 'productos.html').read_text(encoding='utf-8')
cards = re.findall(r'<a href="producto\.html\?id=([^"]+)"[^>]*data-category="([^"]*)"[^>]*data-segment="([^"]*)".*?<img src="([^"]+)" alt="([^"]*)"[^>]*>.*?<h3>([^<]*)</h3>(.*?)</a>', cat, re.DOTALL)
catmeta = {}
for vid, dcat, seg, mainimg, mainalt, name, rest in cards:
    vol = re.search(r'variant-tag">([^<]+)<', rest)
    thumbs = re.findall(r'<img src="([^"]+)"[^>]*title="([^"]*)"', rest)
    cph = {}
    mcolor = mainalt.replace('USWE', '').replace(name, '').strip()
    cph[mcolor.lower()] = mainimg.replace('productos/', '')
    for src, color in thumbs: cph[color.strip().lower()] = src.replace('productos/', '')
    catmeta[vid] = dict(category=dcat, segment=seg, name=name.strip(),
                        vol=vol.group(1).strip() if vol else '', colors=cph,
                        main=mainimg.replace('productos/', ''))

CATLABEL = {'mochila-hidratacion': 'Mochila de Hidratación', 'mochila-race': 'Mochila Race',
            'mochila-running': 'Mochila Running', 'vejiga-hidratacion': 'Vejiga de Hidratación',
            'cinturon-hidratacion': 'Cinturón de Hidratación', 'chaleco-hidratacion': 'Chaleco de Hidratación'}

# id canónico = el mismo id que usa la card del catálogo (match por nombre)
NAME2ID = {m['name'].lower(): vid for vid, m in catmeta.items()}
def cat_id(model): return NAME2ID.get(model.lower(), slug(model))

def resolve_photo(vid, model, color):
    cm = catmeta.get(vid, {})
    # 1) mapa del catalogo por color
    ph = cm.get('colors', {}).get(color.lower())
    if ph and (PRODdir / ph).exists(): return ph
    # 2) por nombre de archivo uswe-<model>-<color>.<ext>
    base = 'uswe-' + slug(model) + '-' + slug(color)
    for f in files:
        if f.lower().rsplit('.',1)[0] == base: return f
    # 3) cualquier variante que empiece con uswe-<model>-<primer token del color>
    ctok = slug(color).split('-')[0]
    for f in files:
        if f.lower().startswith('uswe-'+slug(model)+'-'+ctok): return f
    # 4) fallback: main del catalogo, o alguna del modelo
    if cm.get('main') and (PRODdir/cm['main']).exists(): return cm['main']
    for f in files:
        if f.lower().startswith('uswe-'+slug(model)+'-'): return f
    return None

def describe(name, cont, apto, cat_slug):
    if 'vejiga' in cat_slug or 'bladder' in name.lower():
        return [f"La <em>{name}</em> es una vejiga de hidratación USWE Elite con sistema Plug &amp; Play: se conecta y desconecta en segundos, fácil de limpiar y rellenar.",
                f"{cont}." if cont else "Compatible con las mochilas USWE."]
    return [f"La <em>{name}</em> es una mochila de hidratación USWE con el arnés antivaivén <em>NDM&trade;</em> (No Dancing Monkey), que la mantiene firme al cuerpo sin rebotes en moto, MTB o running.",
            (f"{cont}. " if cont else "") + (f"Apta para {apto.lower()}." if apto else "")]

# --- 2) lista de precios: modelos + colores + specs ---
wb = load_xlsx(SRC, str(Path(r'C:\Users\matia\AppData\Local\Temp\listas') / 'uswe_precios.xlsx')); ws = wb['Precios']
prods = {}     # id -> producto
order = 0
for r in ws.iter_rows(min_row=7, values_only=True):
    if not r[2] or clean(r[1]) != 'USWE': continue
    model = clean(r[3]); color = clean(r[4])
    if 'bladder' in model.lower(): continue          # vejigas: sin stock
    talle = ''
    mp = re.match(r'(Pace \d+)\s*-\s*([ML])$', model)
    if mp:
        if mp.group(2) != 'L': continue              # solo talle L en stock
        model = mp.group(1); talle = 'L'
    vid = cat_id(model)
    cm = catmeta.get(vid, {})
    if vid not in prods:
        order += 1
        cat_slug = cm.get('category', 'mochila-hidratacion')
        prods[vid] = dict(name=model, cat=cat_slug, catlabel=CATLABEL.get(cat_slug, 'Mochila de Hidratación'),
                          seg=cm.get('segment', 'moto'), vol=clean(r[5]), cont=clean(r[6]),
                          apto=clean(r[7]), colors=[], order=order, talle=talle,
                          desc=describe(model, clean(r[6]), clean(r[7]), cat_slug))
    photo = resolve_photo(vid, model, color)
    prods[vid]['colors'].append(dict(color=color, cod=clean(r[2]), disp=clean(r[12]),
                                     photo=('productos/'+photo) if photo else '../../img/logos/uswe.png',
                                     has=bool(photo)))

json.dump(prods, open(r'C:\Users\matia\AppData\Local\Temp\listas\uswe_data.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
tot_colors = sum(len(p['colors']) for p in prods.values())
sinfoto = [(p['name'], c['color']) for p in prods.values() for c in p['colors'] if not c['has']]
print(f"Modelos: {len(prods)} | variantes color: {tot_colors} | sin foto: {len(sinfoto)}")
if sinfoto: print("  Sin foto:", sinfoto)
print("Modelos no en lista (catalogo extra):", [i for i in catmeta if i not in prods])

# ================= FICHA (producto.html) =================
prod = (BASE / 'producto.html').read_text(encoding='utf-8')
PAGEWRAP = '''    <div class="page-wrap">
        <div class="breadcrumb">
            <a href="../../index.html">Inicio</a> <span>&rsaquo;</span>
            <a href="../../productos.html">Productos</a> <span>&rsaquo;</span>
            <a href="productos.html">USWE</a> <span>&rsaquo;</span>
            <span id="bcName">Producto</span>
        </div>
        <div class="product-hero">
            <div class="product-gallery" id="gallery">
                <div class="gallery-main" id="galleryMain">
                    <svg class="gallery-zoom-icon" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">
                        <circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/>
                    </svg>
                    <img id="mainImg" src="" alt="">
                </div>
                <div class="gallery-thumbs" id="galleryThumbs"></div>
            </div>
            <div class="product-info-side">
                <div class="product-kicker" id="pKicker"></div>
                <h1 class="product-title" id="pTitle"></h1>
                <p class="product-tagline" id="pTag"></p>
                <div class="color-pick"><span class="color-pick-label">Color:</span> <span id="pColorName"></span></div>
                <div class="specs-table" id="specsTable"></div>
                <div class="product-description">
                    <div class="description-header">Descripción</div>
                    <div id="pDesc"></div>
                </div>
                <img class="ndm-logo" src="ndm-logo.svg" alt="USWE — No Dancing Monkey">
            </div>
        </div>
    </div>
'''
prod = re.sub(r'    <div class="page-wrap">.*?</div>\s*(?=<footer)', PAGEWRAP, prod, count=1, flags=re.DOTALL)

USWE_CSS = '''        .color-pick { font-size:0.85rem; color:var(--gray); margin-bottom:1.3rem; }
        .color-pick-label { font-weight:700; color:var(--brand-accent); text-transform:uppercase; letter-spacing:1px; font-size:0.7rem; }
        #pColorName { font-weight:600; }
        .gallery-thumb .thumb-color { display:none; }
        .ndm-logo { display:block; width:160px; max-width:55%; margin-top:1.8rem; }
'''
prod = re.sub(r'[ \t]*\.color-pick \{.*?(?=[ \t]*</style>)', '\n', prod, count=1, flags=re.DOTALL)
prod = prod.replace('</style>', USWE_CSS + '</style>', 1)

PAD_SCRIPT = '''    <script>
        const PRODUCTS = ''' + json.dumps(prods, ensure_ascii=False) + ''';
        const id = new URLSearchParams(location.search).get('id');
        const p = PRODUCTS[id] || PRODUCTS[Object.keys(PRODUCTS)[0]];
        const mainImg = document.getElementById('mainImg');
        function setColor(c){
            mainImg.src = c.photo; mainImg.alt = 'USWE ' + p.name + ' ' + c.color;
            document.getElementById('pColorName').textContent = c.color;
            document.getElementById('spColor').textContent = c.color;
            if (document.getElementById('spDisp')) document.getElementById('spDisp').textContent = c.disp || '—';
        }
        if (p) {
            document.title = p.name + ' | USWE · Radsport Geist';
            document.getElementById('bcName').textContent = p.name;
            document.getElementById('pKicker').textContent = p.catlabel;
            document.getElementById('pTitle').textContent = p.name;
            document.getElementById('pTag').textContent = p.apto || '';
            const specs = [];
            if (p.vol) specs.push(['Volumen', p.vol + ' L']);
            specs.push(['Color', '<span id="spColor"></span>']);
            if (p.talle) specs.push(['Talle', p.talle]);
            if (p.cont) specs.push(['Contenidos', p.cont]);
            if (p.apto) specs.push(['Apto para', p.apto]);
            document.getElementById('specsTable').innerHTML = specs.map(function(s){
                return '<div class="specs-row"><span class="specs-key">'+s[0]+'</span><span class="specs-value">'+s[1]+'</span></div>'; }).join('');
            document.getElementById('pDesc').innerHTML = p.desc.map(function(x){ return '<p>'+x+'</p>'; }).join('');
            // thumbs de color
            const tw = document.getElementById('galleryThumbs');
            tw.innerHTML = p.colors.map(function(c,i){
                return '<div class="gallery-thumb'+(i===0?' active':'')+'" data-i="'+i+'" title="'+c.color+'"><img src="'+c.photo+'" alt="'+c.color+'"></div>'; }).join('');
            const th = tw.querySelectorAll('.gallery-thumb');
            th.forEach(function(t){ t.addEventListener('click', function(){
                setColor(p.colors[parseInt(t.dataset.i,10)]);
                th.forEach(function(x){ x.classList.remove('active'); }); t.classList.add('active'); }); });
            setColor(p.colors[0]);
        }
        const galleryMain = document.getElementById('galleryMain');
        if (galleryMain) {
            galleryMain.addEventListener('mousemove', function(e){
                const r = galleryMain.getBoundingClientRect();
                mainImg.style.transformOrigin = ((e.clientX-r.left)/r.width*100)+'% '+((e.clientY-r.top)/r.height*100)+'%';
            });
            galleryMain.addEventListener('mouseleave', function(){ mainImg.style.transformOrigin='center center'; });
        }
    </script>
'''
prod = re.sub(r'    <script>.*?</script>\n', PAD_SCRIPT, prod, count=1, flags=re.DOTALL)
(BASE / 'producto.html').write_text(prod, encoding='utf-8')
print("OK producto.html (ficha USWE con selector de color) generado.")
